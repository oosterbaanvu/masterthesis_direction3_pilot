"""Build a blind review workbook from results/p1_corrections/human_review_batch1.csv.

The workbook contains only the candidate and its documents, never the AI-assisted verdicts, so the
human decision is independent. Odd-numbered items are the required sample (fixed rule, chosen before
looking at any AI verdict per item); even-numbered items are optional.
Output: results/p1_corrections/human_review_batch1_sheet.xlsx
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "results/p1_corrections/human_review_batch1.csv"
OUT = ROOT / "results/p1_corrections/human_review_batch1_sheet.xlsx"

CHANGE = {
    "other_value": "Het getal is veranderd",
    "sign_flip": "Zelfde getal, ander teken (+/-)",
    "retag": "Zelfde getal, ander elementlabel",
    "zero_change": "Van een waarde naar nul, of van nul naar een waarde",
    "period_move": "Zelfde getal, bij een andere periode gerapporteerd",
}
DECISIONS = ["data_entry_error", "legitimate_restatement", "representation_only", "no_difference_found", "unclear"]
WRONG = ["yes", "no", "unknown"]

INSTRUCTIONS = [
    ("Doel", "Per item beoordelen of het oorspronkelijke getal echt fout was, en wat voor soort correctie het is. "
             "Dit is een controle door de student, geen expertoordeel; dat wordt zo gerapporteerd."),
    ("Welke items", "Verplicht: de items met 'ja' in kolom 'verplicht' (de oneven nummers 1, 3, ..., 19). "
                    "De even nummers zijn optioneel."),
    ("Blind blijven", "Open het bestand ai_document_review_batch1_NOT_HUMAN_GOLD.csv NIET voordat je klaar bent."),
    ("Tijd", "Ongeveer 5 tot 10 minuten per item. Lukt het niet binnen 10 minuten: kies 'unclear' en ga door."),
    ("Stap 1", "Open de link 'gecorrigeerd verslag'. Lees bovenaan de 'Explanatory Note': daarin staat waarom het "
               "verslag opnieuw is ingediend."),
    ("Stap 2", "Zoek met Ctrl+F op het getal (kolom 'zoekhint'). Bedragen staan vaak in duizenden: 7,170,767 kan als "
               "7,171 of 7,170,767 staan. Zoek ook op 'restate', 'previously reported' of 'restatement'."),
    ("Stap 3", "Open de link 'origineel verslag' en zoek hetzelfde getal. Vergelijk beide."),
    ("Stap 4", "Kies in kolom 'jouw oordeel' een categorie (keuzelijst) en in 'origineel fout?' ja, nee of unknown."),
    ("Stap 5", "Schrijf in 'waar gevonden' kort de plek, bijv. 'Explanatory Note' of 'Note 2 Restatement tabel', "
               "en in 'toelichting' in een zin waarom. Vul de minuten in."),
    ("Klaar", "Sla op als .xlsx (zelfde naam is prima) en upload het bestand in de chat."),
    ("", ""),
    ("Categorie", "Betekenis"),
    ("data_entry_error", "Het oorspronkelijke getal was een invoer- of tagfout: typefout, verkeerd teken, verkeerde "
                         "schaal (x1000), verkeerd element of verkeerde periode. De onderliggende cijfers waren goed."),
    ("legitimate_restatement", "Het oorspronkelijke getal was boekhoudkundig fout en het bedrijf heeft het herzien, "
                               "bijv. een tabel 'As previously reported / Adjustment / As restated'."),
    ("representation_only", "Dezelfde waarde, alleen anders gepresenteerd of getagd; het oorspronkelijke getal was "
                            "niet echt fout (bijv. een ander maar ook passend element, of een tekenconventie)."),
    ("no_difference_found", "In beide documenten staat hetzelfde getal; het verschil zit waarschijnlijk alleen in de "
                            "data-extractie."),
    ("unclear", "Niet vast te stellen binnen 10 minuten."),
    ("", ""),
    ("origineel fout?", "ja = het oorspronkelijke getal was aantoonbaar fout; nee = het was niet fout; "
                        "unknown = niet vast te stellen."),
    ("Periode", "'periode' = einddatum (jjjjmmdd). 'lengte' = 0 betekent een stand op die datum (balans), "
                "1 = een kwartaal, 2 = een halfjaar, 4 = een heel jaar."),
]


def main() -> None:
    df = pd.read_csv(SRC)
    wb = Workbook()
    info = wb.active
    info.title = "Uitleg"
    for r, (a, b) in enumerate(INSTRUCTIONS, start=1):
        info.cell(r, 1, a).font = Font(bold=a in ("Categorie",) or a.startswith("Stap") or a in ("Doel", "Klaar"))
        info.cell(r, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 110

    ws = wb.create_sheet("Review")
    headers = ["item", "verplicht", "bedrijf", "wat veranderde", "element", "nieuw element", "periode", "lengte",
               "eenheid", "origineel getal", "gecorrigeerd getal", "origineel verslag", "gecorrigeerd verslag",
               "zoekhint", "jouw oordeel", "origineel fout?", "waar gevonden", "toelichting", "minuten"]
    ws.append(headers)
    for _, r in df.iterrows():
        ws.append([int(r["item"]), "ja" if int(r["item"]) % 2 == 1 else "optioneel", r["name"],
                   CHANGE.get(r["correction_type"], r["correction_type"]), r["tag"],
                   "" if pd.isna(r["new_tag"]) else r["new_tag"], str(int(r["period_end"])), int(r["qtrs"]),
                   r["uom"], r["orig_value"], r["amend_value"], r["orig_document"], r["amend_document"],
                   r["search_hint"], None, None, None, None, None])
    n = len(df) + 1
    for row in range(2, n + 1):
        for col, label in ((12, "origineel"), (13, "gecorrigeerd")):
            c = ws.cell(row, col)
            c.hyperlink = c.value
            c.value = f"open {label} verslag"
            c.font = Font(color="1F5FBF", underline="single")
        for col in (10, 11):
            ws.cell(row, col).number_format = "#,##0"
        fill = PatternFill("solid", fgColor="FFF7D6")
        for col in range(15, 20):
            ws.cell(row, col).fill = fill
    for col in range(1, len(headers) + 1):
        ws.cell(1, col).font = Font(bold=True)
    dv1 = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(WRONG) + '"', allow_blank=True)
    dv3 = DataValidation(type="whole", operator="between", formula1="0", formula2="120", allow_blank=True)
    ws.add_data_validation(dv1); ws.add_data_validation(dv2); ws.add_data_validation(dv3)
    dv1.add(f"O2:O{n}"); dv2.add(f"P2:P{n}"); dv3.add(f"S2:S{n}")
    widths = [6, 10, 30, 34, 40, 26, 10, 7, 8, 14, 14, 20, 22, 46, 22, 14, 26, 50, 8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "D2"
    ws.cell(n + 2, 1, "Aantal ingevuld (verplicht):")
    ws.cell(n + 2, 3, f'=COUNTIFS(B2:B{n},"ja",O2:O{n},"<>")')
    ws.cell(n + 3, 1, "Totaal minuten:")
    ws.cell(n + 3, 3, f"=SUM(S2:S{n})")
    wb.active = 0
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
